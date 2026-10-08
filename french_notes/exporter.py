"""Export saved French notes to a local Word document."""

from __future__ import annotations

import os
from pathlib import Path
import tempfile
from typing import TYPE_CHECKING, Sequence

if TYPE_CHECKING:
    from .core import Note


_DOCUMENT_TEXT = {
    "fr": {
        "locale": "fr-FR",
        "title": "Notes de français",
        "words": "Mots",
        "sentences": "Phrases",
        "empty_words": "Aucun mot enregistré.",
        "empty_sentences": "Aucune phrase enregistrée.",
        "extension_error": "Le fichier Word exporté doit porter l’extension .docx.",
        "empty_error": "Aucune note à exporter.",
        "category_error": "Chaque note doit appartenir à la catégorie « mot » ou « phrase ».",
    },
    "en": {
        "locale": "en-US",
        "title": "French Notes",
        "words": "Words",
        "sentences": "Sentences",
        "empty_words": "No saved words.",
        "empty_sentences": "No saved sentences.",
        "extension_error": "The exported Word file must use the .docx extension.",
        "empty_error": "There are no notes to export.",
        "category_error": "Each note must belong to the word or sentence category.",
    },
    "zh": {
        "locale": "zh-CN",
        "title": "法语笔记",
        "words": "单词",
        "sentences": "句子",
        "empty_words": "暂无单词记录。",
        "empty_sentences": "暂无句子记录。",
        "extension_error": "导出的 Word 文件必须使用 .docx 扩展名。",
        "empty_error": "没有可导出的笔记。",
        "category_error": "每条笔记必须归入单词或句子类别。",
    },
}


def _summary(language: str, total: int, words: int, sentences: int) -> str:
    if language == "zh":
        return (
            f"共 {total} 条笔记，其中单词 {words} 条，句子 {sentences} 条。"
            "以下按类别列出已保存的法语内容。"
        )
    if language == "en":
        note_label = "saved note" if total == 1 else "saved notes"
        word_label = "word" if words == 1 else "words"
        sentence_label = "sentence" if sentences == 1 else "sentences"
        return (
            f"{total} {note_label}: {words} {word_label} and {sentences} {sentence_label}. "
            "Notes are grouped by category."
        )
    note_label = "note enregistrée" if total == 1 else "notes enregistrées"
    word_label = "mot" if words < 2 else "mots"
    sentence_label = "phrase" if sentences < 2 else "phrases"
    return (
        f"{total} {note_label} : {words} {word_label} et {sentences} {sentence_label}. "
        "Le contenu est regroupé par catégorie."
    )


def export_docx(notes: Sequence[Note], path: Path, language: str = "fr") -> Path:
    """Export all notes, grouped by category, without risking an existing file.

    ``python-docx`` is imported only when exporting. Saving uses a temporary
    file in the destination directory, so replacing the completed document is
    atomic. An open Word document on Windows can prevent replacement; that
    error is propagated while the original document is retained.

    ``language`` selects document labels and validation messages (``fr``,
    ``en`` or ``zh``), while the stored French note text remains unchanged.
    French remains the default for existing two-argument callers.
    """
    if language not in ("fr", "en", "zh"):
        raise ValueError("Unsupported document language. Choose fr, en or zh.")
    text = _DOCUMENT_TEXT[language]
    path = Path(path)
    if path.suffix.lower() != ".docx":
        raise ValueError(text["extension_error"])
    records = tuple(notes)
    if not records:
        raise ValueError(text["empty_error"])
    if any(note.category not in {"word", "sentence"} for note in records):
        raise ValueError(text["category_error"])

    # Keep CSV reading and deduplication usable without the optional exporter.
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)

    for style_name in ("Normal", "Title", "Heading 1"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        run_properties = style.element.get_or_add_rPr()
        fonts = run_properties.get_or_add_rFonts()
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        # Theme formatting must not override the deliberately chosen fonts.
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        language_element = run_properties.find(qn("w:lang"))
        if language_element is None:
            language_element = OxmlElement("w:lang")
            run_properties.append(language_element)
        language_element.set(qn("w:val"), text["locale"])
        color = run_properties.find(qn("w:color"))
        if color is not None:
            for attribute in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn(f"w:{attribute}"), None)
        paragraph_properties = style.element.find(qn("w:pPr"))
        if paragraph_properties is not None:
            border = paragraph_properties.find(qn("w:pBdr"))
            if border is not None:
                paragraph_properties.remove(border)

    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.widow_control = True
    document.styles["Title"].font.size = Pt(24)
    document.styles["Title"].paragraph_format.space_after = Pt(12)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(6)

    words = [note for note in records if note.category == "word"]
    sentences = [note for note in records if note.category == "sentence"]
    groups = (
        (text["words"], words, text["empty_words"]),
        (text["sentences"], sentences, text["empty_sentences"]),
    )
    document.add_paragraph(text["title"], style="Title")
    document.add_paragraph(_summary(language, len(records), len(words), len(sentences)))
    for heading, group, empty_message in groups:
        document.add_heading(heading, level=1)
        if not group:
            document.add_paragraph(empty_message)
            continue
        for number, note in enumerate(group, start=1):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.first_line_indent = Cm(-0.8)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))
            # Long sentences may flow across pages instead of leaving gaps.
            paragraph.paragraph_format.keep_together = False
            paragraph.add_run(f"{number}.\t")
            content = paragraph.add_run(note.french)
            french_language = OxmlElement("w:lang")
            french_language.set(qn("w:val"), "fr-FR")
            content._r.get_or_add_rPr().append(french_language)

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as destination:
            document.save(destination)
            destination.flush()
            os.fsync(destination.fileno())
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return path
