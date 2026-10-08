"""Run real application checks in a new, isolated directory.

This entry point is also used by the packaged Windows executable.  It exercises
Tk and Word export without opening dialogs or accessing the user's note store.
"""

from __future__ import annotations

import json
from pathlib import Path
import sys
import tempfile
import traceback
import unicodedata
from typing import Any

from . import __version__
from .core import CSVStore, normalize_key
from .exporter import export_docx


REPORT_NAME = "self-test-result.json"


def _require(condition: bool, message: str) -> None:
    # Do not use assert: the packaged executable may run optimized Python.
    if not condition:
        raise RuntimeError(message)


def _diagnostic_folder(preferred_parent: Path) -> Path:
    """Reserve a fresh location when the requested directory is unsafe."""
    try:
        return Path(tempfile.mkdtemp(prefix="french-notes-self-test-error-", dir=preferred_parent))
    except OSError:
        return Path(tempfile.mkdtemp(prefix="french-notes-self-test-error-"))


def _announce(report_path: Path, success: bool) -> None:
    # PyInstaller --windowed sets stdout/stderr to None on Windows.
    stream = sys.stdout if success else sys.stderr
    if stream is not None:
        print(f"FrenchNotes self-test {'passed' if success else 'failed'}: {report_path}", file=stream)


def run_self_test(folder: Path) -> int:
    """Exercise CSV, the real Tk interface and DOCX; return a process exit code.

    ``folder`` must not already exist, even if it is empty.  Exclusive directory
    creation prevents a test from modifying a user's CSV or previous test run.
    A conflict report goes into a separately reserved diagnostic directory,
    whose path is printed when a console stream is available.
    """
    folder = Path(folder).expanduser().absolute()
    result: dict[str, Any] = {"success": False, "version": __version__, "checks": []}
    owned_folder = False
    root = None
    try:
        folder.mkdir(parents=True, exist_ok=False)
        owned_folder = True
        result["checks"].append("isolated_directory")

        csv_path = folder / "notes.csv"
        store = CSVStore(csv_path)
        first_batch = "École\nécole\nJe suis étudiant.\nà\na"
        preview = store.preview(first_batch)
        _require(
            [candidate.status for candidate in preview] == ["new", "batch", "new", "new", "new"],
            "The first batch must contain four new notes and one batch duplicate.",
        )
        _require(
            [candidate.category for candidate in preview] == ["word", "word", "sentence", "word", "word"],
            "Word/sentence classification differs from the expected categories.",
        )
        _require(not csv_path.exists(), "Preview unexpectedly wrote a CSV.")
        saved = store.append(preview)
        _require(len(saved.added) == 4 and saved.skipped == 1, "Save must add four notes and skip one.")
        result["checks"].append("csv_preview_classify_and_save")

        _require(normalize_key("E\u0301cole") == normalize_key("École"), "NFC accents were not normalized.")
        _require(normalize_key("à") != normalize_key("a"), "Accented and unaccented words must remain distinct.")
        _require(normalize_key("Je suis étudiant.") != normalize_key("Je suis étudiant"), "Punctuation differences must remain distinct.")
        notes = CSVStore(csv_path).load()
        _require([note.french for note in notes] == ["École", "Je suis étudiant.", "à", "a"], "Reloaded CSV lost or changed French text.")
        _require(all(unicodedata.is_normalized("NFC", note.french) for note in notes), "Saved French text is not NFC.")
        result["checks"].append("csv_reload_and_unicode")

        csv_before = csv_path.read_bytes()
        duplicates = store.preview(" e\u0301COLE \n JE   SUIS étudiant. \n À \n A")
        _require(all(candidate.status == "existing" for candidate in duplicates), "Case/whitespace/NFC duplicates were not recognized.")
        duplicate_result = store.append(duplicates)
        _require(not duplicate_result.added and duplicate_result.skipped == 4, "Duplicate batch was unexpectedly saved.")
        _require(csv_path.read_bytes() == csv_before, "Duplicate-only input changed the CSV bytes.")
        _require(not store.backup_path.exists(), "Duplicate-only input unexpectedly created a backup.")
        result["checks"].append("duplicate_batch_preserves_csv")

        # Import the actual modules here so a missing bundled dependency becomes
        # a JSON failure rather than a startup crash before a report is written.
        import tkinter as tk
        from .ui import FrenchNotesApp

        class SelfTestApp(FrenchNotesApp):
            def _read_settings(self) -> dict:
                # Self-test must not read or change the user's preferences.
                return {}

            def _error(self, title: str, error: Exception) -> None:
                # The production method opens a modal dialog.  Convert the same
                # error to an exception so an unattended build cannot hang.
                raise RuntimeError(f"{title}: {error}") from error

        callback_errors: list[BaseException] = []
        root = tk.Tk()
        root.withdraw()
        root.report_callback_exception = lambda _kind, error, _tb: callback_errors.append(error)
        app = SelfTestApp(root, csv_path=csv_path, language="fr")

        def update_root() -> None:
            root.update()
            if callback_errors:
                raise RuntimeError("Tk callback failed during self-test.") from callback_errors[0]

        update_root()
        _require(app.store.path == csv_path, "The interface did not use the explicit test CSV.")
        _require(len(app.notes) == 4 and len(app.library_tree.get_children()) == 4, "The interface did not load the saved library.")
        _require("Notes de français" in root.title(), "The interface window title is not in French.")
        categories = {app.library_tree.item(item, "values")[0] for item in app.library_tree.get_children()}
        _require(categories == {"Mot", "Phrase"}, "The library categories are not displayed in French.")
        _require(root.state() == "withdrawn", "Self-test exposed its application window.")
        result["checks"].append("tk_interface_and_existing_library")

        app.input.insert("1.0", "ÉCOLE\npomme de terre\nOù est le café ?")
        update_root()
        app.analyze()
        update_root()
        _require([candidate.status for candidate in app.candidates] == ["existing", "new", "new"], "UI analysis did not identify existing/new content.")
        _require(app.candidates[1].category == "sentence", "The phrase did not receive its default classification.")
        app.preview_tree.selection_set("1")
        app.set_category("word")
        update_root()
        _require(app.candidates[1].category == "word", "Manual UI category correction failed.")
        from .i18n import BRAND

        input_before = app.input.get("1.0", "end-1c")
        candidate_before = [(item.french, item.category, item.status) for item in app.candidates]
        for language, title in (("zh", "法语笔记"), ("en", "French Notes"), ("fr", "Notes de français")):
            app.set_language(language, persist=False)
            update_root()
            _require(title in root.title(), "Language switch did not translate the window title.")
            _require(app.input.get("1.0", "end-1c") == input_before, "Language switch changed unsaved input.")
            _require([(item.french, item.category, item.status) for item in app.candidates] == candidate_before, "Language switch changed preview/category state.")
            _require(app.preview_tree.selection() == ("1",), "Language switch lost the selected candidate.")
            _require("disabled" not in app.save_button.state(), "Language switch disabled a valid save.")
            _require(app.brand_label.cget("text") == BRAND, "The company watermark is missing.")
        result["checks"].append("tk_language_switch_preserves_input")
        # Use the real store directly to avoid save dialogs even on unusual
        # filesystem failures (for example, lock cleanup failure).
        ui_saved = app.store.append(app.candidates)
        _require(len(ui_saved.added) == 2 and ui_saved.skipped == 1, "The analyzed UI batch was not saved correctly.")
        app.refresh_library()
        update_root()
        notes = CSVStore(csv_path).load()
        _require(len(notes) == 6 and len(app.library_tree.get_children()) == 6, "The saved UI batch was not reloaded.")
        _require(next(note for note in notes if note.french == "pomme de terre").category == "word", "Manual category was not persisted.")
        _require(not store.backup_path.exists(), "A successful update retained an old CSV backup.")
        result["checks"].append("tk_analysis_and_manual_category")

        output_path = export_docx(notes, folder / "notes.docx")
        from docx import Document

        parsed_document = Document(output_path)
        paragraphs = [paragraph.text for paragraph in parsed_document.paragraphs]
        exported_notes = [paragraph.split("\t", 1)[1] for paragraph in paragraphs if "\t" in paragraph]
        _require(sorted(exported_notes) == sorted(note.french for note in notes), "Word export did not preserve every French note exactly once.")
        _require("Notes de français" in paragraphs, "Le titre de l’export Word est absent.")
        _require("Mots" in paragraphs and "Phrases" in paragraphs, "Les titres des catégories sont absents de l’export Word.")
        result["checks"].append("docx_export_and_parse")
        for language, title, headings in (("en", "French Notes", ("Words", "Sentences")), ("zh", "法语笔记", ("单词", "句子"))):
            translated_path = export_docx(notes, folder / f"notes-{language}.docx", language=language)
            translated = [paragraph.text for paragraph in Document(translated_path).paragraphs]
            _require(title in translated and all(heading in translated for heading in headings), "Word export labels did not follow the selected language.")
            translated_notes = [paragraph.split("\t", 1)[1] for paragraph in translated if "\t" in paragraph]
            _require(sorted(translated_notes) == sorted(note.french for note in notes), "A translated Word export changed the French notes.")
        result["checks"].append("docx_multilingual_labels")
        result["success"] = True
    except Exception as error:
        result["error"] = str(error)
        result["traceback"] = traceback.format_exc()
    finally:
        if root is not None:
            try:
                root.destroy()
            except Exception as error:
                if result["success"]:
                    result["success"] = False
                    result["error"] = f"Could not destroy self-test Tk root: {error}"
                    result["traceback"] = traceback.format_exc()

    report_folder = folder if owned_folder else _diagnostic_folder(folder.parent)
    report_path = report_folder / REPORT_NAME
    try:
        # Never replace an existing report, including one introduced while the
        # test was running.  Diagnostics also use an exclusively created path.
        with report_path.open("x", encoding="utf-8") as report:
            json.dump(result, report, ensure_ascii=False, indent=2)
            report.write("\n")
    except OSError as error:
        result["success"] = False
        result["error"] = f"Could not write self-test report: {error}"
        result["traceback"] = traceback.format_exc()
        report_path = _diagnostic_folder(folder.parent) / REPORT_NAME
        with report_path.open("x", encoding="utf-8") as report:
            json.dump(result, report, ensure_ascii=False, indent=2)
            report.write("\n")
    _announce(report_path, result["success"])
    return 0 if result["success"] else 1
