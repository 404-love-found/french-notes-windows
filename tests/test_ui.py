"""End-to-end desktop flows against a temporary local CSV."""

import json
from pathlib import Path
import sys
import tempfile
import tkinter as tk
import unittest
from unittest.mock import patch

from french_notes.core import CSVStore
from french_notes.i18n import BRAND, LANGUAGES, translate
from french_notes.ui import FrenchNotesApp


class DesktopFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.path = self.folder / "notes.csv"
        try:
            self.root = tk.Tk()
        except tk.TclError as error:
            if sys.platform == "win32":
                raise
            self.skipTest(f"Desktop display unavailable: {error}")
        self.root.withdraw()
        self.addCleanup(self.root.destroy)
        self.callback_errors = []
        self.root.report_callback_exception = self.record_callback_error
        self.settings_patch = patch("french_notes.ui.app_data_dir", return_value=self.folder / "settings")
        self.settings_patch.start()
        self.addCleanup(self.settings_patch.stop)
        self.messages = patch("french_notes.ui.messagebox.showerror")
        self.error_dialog = self.messages.start()
        self.addCleanup(self.messages.stop)
        self.app = FrenchNotesApp(self.root, self.path)
        self.root.update()

    def record_callback_error(self, error_type, error, _traceback):
        self.callback_errors.append(f"{error_type.__name__}: {error}")

    def tearDown(self):
        self.assertEqual(self.callback_errors, [], "Tk callbacks must not fail silently")

    def configured_app(self, **options):
        """Create another isolated view to exercise settings reads on restart."""
        window = tk.Toplevel(self.root)
        window.withdraw()
        self.addCleanup(window.destroy)
        app = FrenchNotesApp(window, **options)
        self.root.update()
        return app

    def enter(self, text):
        self.app.input.delete("1.0", "end")
        self.app.input.insert("1.0", text)
        self.root.update()
        self.app.analyze()
        self.root.update()

    def test_input_preview_save_reload_and_search(self):
        self.enter("École\nécole\nJe suis étudiant.\n\nà")
        self.assertEqual([item.status for item in self.app.candidates], ["new", "batch", "new", "new"])
        self.assertFalse(self.path.exists(), "Preview must not write the CSV")
        self.app.save_new()
        self.root.update()
        self.assertEqual(self.app.input.get("1.0", "end-1c"), "")
        self.assertEqual(len(self.app.store.load()), 3)
        self.assertIn("Notes enregistrées : 3", self.app.summary.get())
        self.app.search.set("ÉCOLE")
        self.assertEqual(len(self.app.library_tree.get_children()), 1)
        self.app.search.set("")
        self.app.category_filter.set("Phrases")
        self.assertEqual(len(self.app.library_tree.get_children()), 1)
        self.app.category_filter.set("Mots")
        self.assertEqual(len(self.app.library_tree.get_children()), 2)
        self.app.category_filter.set("Tous")
        self.assertEqual(len(self.app.library_tree.get_children()), 3)
        self.enter(" école \nJE  SUIS étudiant.")
        self.assertEqual([item.status for item in self.app.candidates], ["existing", "existing"])
        self.assertIn("disabled", self.app.save_button.state())
        self.assertFalse(self.path.with_suffix(".csv.bak").exists())

    def test_editing_input_invalidates_previous_preview(self):
        self.enter("bonjour")
        self.assertNotIn("disabled", self.app.save_button.state())
        self.app.input.insert("end", "\nBonsoir")
        self.root.update()
        self.assertEqual(self.app.candidates, [])
        self.assertIn("disabled", self.app.save_button.state())
        self.assertFalse(self.path.exists())

    def test_save_and_export_buttons_fit_default_and_minimum_windows(self):
        # Windows does not lay out a withdrawn toplevel. Map a transparent test
        # window so the actual geometry is calculated without obscuring the desktop.
        self.root.attributes("-alpha", 0.0)
        self.root.deiconify()
        self.root.update()

        def check_bounds():
            self.app.notebook.select(self.app.intake_tab)
            for language in ("fr", "en", "zh"):
                self.app.set_language(language, persist=False)
                for width, height in ((1180, 760), (1024, 640)):
                    self.root.geometry(f"{width}x{height}")
                    self.root.update()
                    actual_size = (self.root.winfo_width(), self.root.winfo_height())
                    if (width, height) == (1024, 640):
                        self.assertEqual(actual_size, (width, height))
                    else:
                        # A small desktop can constrain the default window. Check
                        # controls in its actual viewport, including that constraint.
                        self.assertGreaterEqual(actual_size[0], 1024)
                        self.assertGreaterEqual(actual_size[1], 640)
                    controls = (
                        ("save", self.app.save_button),
                        ("export", self.app.export_button),
                        ("language_selector", self.app.language_selector),
                        ("brand", self.app.brand_label),
                    )
                    for name, widget in controls:
                        with self.subTest(
                            language=language, requested=(width, height), actual=actual_size,
                            control=name, save_enabled="disabled" not in self.app.save_button.state(),
                        ):
                            # Do not filter unmapped widgets: that would silently
                            # omit controls hidden by a layout regression.
                            self.assertTrue(widget.winfo_ismapped())
                            self.assertGreaterEqual(widget.winfo_width(), widget.winfo_reqwidth())
                            self.assertGreaterEqual(widget.winfo_height(), widget.winfo_reqheight())
                            x = widget.winfo_rootx() - self.root.winfo_rootx()
                            y = widget.winfo_rooty() - self.root.winfo_rooty()
                            self.assertGreaterEqual(x, 0)
                            self.assertGreaterEqual(y, 0)
                            self.assertLessEqual(x + widget.winfo_width(), self.root.winfo_width())
                            self.assertLessEqual(y + widget.winfo_height(), self.root.winfo_height())
                    self.assertEqual(self.app.brand_label.cget("text"), BRAND)

        check_bounds()
        self.enter("bonjour\nBonjour\nJe suis étudiant.")
        self.assertNotIn("disabled", self.app.save_button.state())
        check_bounds()
        self.app.save_new()
        self.root.update()
        self.assertNotIn("disabled", self.app.export_button.state())
        check_bounds()
        self.enter("bonjour\nBonjour\nJe suis étudiant.")
        self.assertIn("disabled", self.app.save_button.state())
        check_bounds()

    def test_manual_category_is_saved(self):
        self.enter("pomme de terre")
        self.app.preview_tree.selection_set("0")
        self.app.set_category("word")
        self.app.save_new()
        self.root.update()
        self.assertEqual(self.app.store.load()[0].category, "word")

    def test_save_failure_preserves_unsaved_text(self):
        self.enter("aujourd’hui")
        with patch.object(self.app.store, "append", side_effect=OSError("disk is full")):
            self.app.save_new()
        self.assertEqual(self.app.input.get("1.0", "end-1c"), "aujourd’hui")
        self.assertEqual(len(self.app.candidates), 1)
        self.assertFalse(self.path.exists())
        self.error_dialog.assert_called_once()

    def test_storage_error_keeps_input_and_shows_dialog(self):
        self.enter("bonjour")
        self.app.store.lock_path.write_text("another instance", encoding="utf-8")
        self.app.save_new()
        self.assertEqual(self.app.input.get("1.0", "end-1c"), "bonjour")
        self.assertFalse(self.path.exists())
        self.error_dialog.assert_called_once()

    def test_bad_csv_switch_does_not_replace_current_path(self):
        self.enter("bonjour")
        malformed = self.folder / "invalid.csv"
        malformed.write_text("french,category\nbonjour,word\n", encoding="utf-8")
        with patch("french_notes.ui.filedialog.asksaveasfilename", return_value=str(malformed)):
            self.app.choose_csv()
        self.assertEqual(self.app.store.path, self.path)
        self.assertEqual(self.app.input.get("1.0", "end-1c"), "bonjour")
        self.error_dialog.assert_called_once()

    def test_saved_data_with_lock_cleanup_error_is_reported_as_saved(self):
        self.enter("école")
        original_unlink = Path.unlink

        def fail_only_lock(path, *args, **kwargs):
            if path == self.app.store.lock_path:
                raise PermissionError("lock cannot be removed")
            return original_unlink(path, *args, **kwargs)

        with patch.object(Path, "unlink", autospec=True, side_effect=fail_only_lock), patch("french_notes.ui.messagebox.showwarning") as warning:
            self.app.save_new()
        self.root.update()
        self.assertEqual(self.app.input.get("1.0", "end-1c"), "")
        self.assertEqual(len(self.app.store.load()), 1)
        self.assertEqual(len(self.app.notes), 1)
        self.assertIn("Notes enregistrées : 1", self.app.summary.get())
        warning.assert_called_once()
        self.error_dialog.assert_not_called()

    def test_switch_csv_keeps_input_and_requires_new_analysis(self):
        self.enter("bonjour")
        replacement = self.folder / "other.csv"
        with patch("french_notes.ui.filedialog.asksaveasfilename", return_value=str(replacement)):
            self.app.choose_csv()
        self.assertEqual(self.app.store.path, replacement)
        self.assertEqual(self.app.input.get("1.0", "end-1c"), "bonjour")
        self.assertEqual(self.app.candidates, [])
        self.assertIn("disabled", self.app.save_button.state())
        self.app.analyze()
        self.app.save_new()
        self.assertTrue(replacement.exists())
        self.assertFalse(self.path.exists())

    def test_export_contains_entire_csv_despite_filter(self):
        from docx import Document

        self.enter("école\nJe suis étudiant.")
        self.app.save_new()
        self.root.update()
        self.app.search.set("école")
        output = self.folder / "notes.docx"
        with patch("french_notes.ui.filedialog.asksaveasfilename", return_value=str(output)), patch("french_notes.ui.messagebox.showinfo"):
            self.app.export_all()
        text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
        self.assertIn("école", text)
        self.assertIn("Je suis étudiant.", text)

    def test_language_switch_preserves_unsaved_notes_categories_filters_and_selection(self):
        self.enter("école\nJe suis étudiant.")
        self.app.save_new()
        self.root.update()
        original_csv = self.path.read_bytes()
        original_mtime = self.path.stat().st_mtime_ns
        saved_notes = list(self.app.notes)
        unsaved = "bonjour\nBonjour\nécole\npomme de terre"
        self.enter(unsaved)
        self.app.preview_tree.selection_set("3")
        self.app.set_category("word")
        self.app.preview_tree.selection_set(("0", "3"))
        self.app.category_filter.set("Phrases")
        self.app.search.set("étudiant")
        self.app.library_tree.selection_set("1")
        expected_candidates = [(id(item), item.french, item.category, item.status) for item in self.app.candidates]
        preview_selection = self.app.preview_tree.selection()
        library_selection = self.app.library_tree.selection()

        for language in ("en", "zh", "fr"):
            with self.subTest(language=language):
                self.app.set_language(language, persist=False)
                self.root.update()
                self.assertEqual(self.app.language, language)
                self.assertEqual(self.app.language_selector.get(), LANGUAGES[language])
                self.assertEqual(self.app.input.get("1.0", "end-1c"), unsaved)
                self.assertEqual(
                    [(id(item), item.french, item.category, item.status) for item in self.app.candidates],
                    expected_candidates,
                )
                self.assertEqual(self.app.notes, saved_notes)
                self.assertEqual(self.app.preview_tree.selection(), preview_selection)
                self.assertEqual(self.app.library_tree.selection(), library_selection)
                self.assertEqual(self.app.search.get(), "étudiant")
                self.assertEqual(self.app._filters()[self.app.category_filter.get()], "sentence")
                self.assertEqual(self.app.library_tree.get_children(), ("1",))
                self.assertEqual(self.app.library_tree.item("1", "values")[1], "Je suis étudiant.")
                self.assertNotIn("disabled", self.app.save_button.state())
                self.assertNotIn("disabled", self.app.export_button.state())
                self.assertEqual(
                    self.app.summary.get(),
                    translate(language, "preview_summary", new=2, existing=1, batch=1),
                )
                self.assertEqual(
                    self.app.status.get(),
                    translate(language, "reclassified", changed=1, category=translate(language, "word")),
                )
                for row_id, item in zip(self.app.preview_tree.get_children(), self.app.candidates):
                    self.assertEqual(
                        self.app.preview_tree.item(row_id, "values"),
                        (translate(language, item.category), item.french, translate(language, f"status_{item.status}")),
                    )
                self.assertEqual(self.app.preview_tree.heading("french", "text"), translate(language, "content"))
                self.assertEqual(self.app.library_tree.heading("created", "text"), translate(language, "created"))
                self.assertEqual(self.app.brand_label.cget("text"), BRAND)
                self.assertEqual(self.path.read_bytes(), original_csv)
                self.assertEqual(self.path.stat().st_mtime_ns, original_mtime)
        self.assertFalse(self.app.config_path.exists(), "QA language switching must not persist settings")
        self.error_dialog.assert_not_called()

    def test_language_switch_preserves_empty_and_duplicate_only_states(self):
        for language in ("en", "zh", "fr"):
            with self.subTest(state="empty", language=language):
                self.app.set_language(language, persist=False)
                self.root.update()
                self.assertEqual(self.app.status.get(), translate(language, "ready"))
                self.assertEqual(self.app.candidates, [])
                self.assertIn("disabled", self.app.save_button.state())
                self.assertIn("disabled", self.app.export_button.state())
                self.assertFalse(self.path.exists())
        self.enter("bonjour")
        self.app.save_new()
        self.root.update()
        original = self.path.read_bytes()
        self.enter("BONJOUR\nbonjour")
        for language in ("en", "zh", "fr"):
            with self.subTest(state="duplicates", language=language):
                self.app.set_language(language, persist=False)
                self.root.update()
                self.assertEqual([item.status for item in self.app.candidates], ["existing", "existing"])
                self.assertEqual(self.app.input.get("1.0", "end-1c"), "BONJOUR\nbonjour")
                self.assertIn("disabled", self.app.save_button.state())
                self.assertNotIn("disabled", self.app.export_button.state())
                self.assertEqual(self.app.status.get(), translate(language, "no_new"))
                self.assertEqual(self.app.summary.get(), translate(language, "preview_summary", new=0, existing=2, batch=0))
                self.assertEqual(self.path.read_bytes(), original)

    def test_missing_corrupt_and_unknown_language_settings_fall_back_to_french(self):
        config = self.app.config_path
        config.parent.mkdir(parents=True, exist_ok=True)
        cases = (
            ("missing", None),
            ("invalid_json", '{"language":'),
            ("wrong_settings_type", '["en"]'),
            ("unknown_language", json.dumps({"language": "unknown"})),
            ("wrong_language_type", json.dumps({"language": ["en"]})),
        )
        for case, content in cases:
            with self.subTest(config=case):
                if content is None:
                    config.unlink(missing_ok=True)
                else:
                    config.write_text(content, encoding="utf-8")
                restored = self.configured_app()
                self.assertEqual(restored.language, "fr")
                self.assertEqual(restored.language_selector.get(), "Français")
                self.assertEqual(restored.store.path, config.parent / "notes.csv")
                self.assertEqual(restored.notes, [])
                self.assertIn("Notes de français", restored.root.title())
        self.error_dialog.assert_not_called()

    def test_language_and_csv_path_persist_and_reload_without_rewriting_notes(self):
        self.enter("école\nJe suis étudiant.")
        self.app.save_new()
        self.root.update()
        original_csv = self.path.read_bytes()
        original_mtime = self.path.stat().st_mtime_ns
        config = self.app.config_path
        config.parent.mkdir(parents=True, exist_ok=True)
        config.write_text(
            json.dumps({"language": "fr", "csv_path": str(self.path), "keep_setting": {"value": 7}}),
            encoding="utf-8",
        )
        restored = self.configured_app()
        self.assertEqual(restored.store.path, self.path)
        restored.set_language("en")
        self.root.update()
        saved = json.loads(config.read_text(encoding="utf-8"))
        self.assertEqual(saved["language"], "en")
        self.assertEqual(saved["csv_path"], str(self.path))
        self.assertEqual(saved["keep_setting"], {"value": 7})
        reloaded = self.configured_app()
        self.assertEqual(reloaded.language, "en")
        self.assertEqual(reloaded.store.path, self.path)
        self.assertEqual(reloaded.notes, self.app.notes)

        # Exercise the real ComboboxSelected binding, not only the public setter.
        reloaded.language_selector.set(LANGUAGES["zh"])
        reloaded.language_selector.event_generate("<<ComboboxSelected>>")
        self.root.update()
        self.assertEqual(reloaded.language, "zh")
        saved = json.loads(config.read_text(encoding="utf-8"))
        self.assertEqual(saved["language"], "zh")
        self.assertEqual(saved["csv_path"], str(self.path))
        self.assertEqual(saved["keep_setting"], {"value": 7})
        final_reload = self.configured_app()
        self.assertEqual(final_reload.language, "zh")
        self.assertEqual(final_reload.notes, self.app.notes)
        self.assertEqual(self.path.read_bytes(), original_csv)
        self.assertEqual(self.path.stat().st_mtime_ns, original_mtime)
        self.assertFalse(config.with_suffix(".tmp").exists())

    def test_unknown_language_fallback_keeps_the_configured_csv_path(self):
        self.enter("bonjour")
        self.app.save_new()
        config = self.app.config_path
        config.parent.mkdir(parents=True, exist_ok=True)
        config.write_text(json.dumps({"language": "unknown", "csv_path": str(self.path)}), encoding="utf-8")
        original = self.path.read_bytes()
        restored = self.configured_app()
        self.assertEqual(restored.language, "fr")
        self.assertEqual(restored.store.path, self.path)
        self.assertEqual([note.french for note in restored.notes], ["bonjour"])
        self.assertEqual(self.path.read_bytes(), original)

    def test_switching_csv_does_not_delete_other_csv_backup_or_word_files(self):
        self.enter("bonjour")
        self.app.save_new()
        replacement = self.folder / "selected.csv"
        store = CSVStore(replacement)
        store.append(store.preview("merci"))
        self.path.with_suffix(".csv.bak").write_bytes(b"keep the previous backup")
        archive = self.folder / "archive.csv"
        archive.write_bytes(b"keep unrelated historical data")
        document = self.folder / "saved-notes.docx"
        document.write_bytes(b"keep the existing Word document")
        protected = (self.path, self.path.with_suffix(".csv.bak"), archive, document, replacement)
        before = {path: path.read_bytes() for path in protected}
        self.enter("une note non enregistrée")
        self.app.set_language("zh", persist=False)
        with patch("french_notes.ui.filedialog.asksaveasfilename", return_value=str(replacement)):
            self.app.choose_csv()
        self.root.update()
        self.assertEqual(self.app.store.path, replacement)
        self.assertEqual(self.app.input.get("1.0", "end-1c"), "une note non enregistrée")
        self.assertEqual([note.french for note in self.app.notes], ["merci"])
        self.assertEqual(self.app.candidates, [])
        self.assertIn("disabled", self.app.save_button.state())
        self.assertEqual(self.app.status.get(), translate("zh", "csv_changed"))
        for path, original in before.items():
            with self.subTest(protected_file=path.name):
                self.assertEqual(path.read_bytes(), original)
        settings = json.loads(self.app.config_path.read_text(encoding="utf-8"))
        self.assertEqual(settings["csv_path"], str(replacement))
        self.assertEqual(settings["language"], "zh")

    def test_export_document_and_dialog_follow_each_interface_language(self):
        from docx import Document

        self.enter("école\nJe suis étudiant.")
        self.app.save_new()
        self.root.update()
        original_csv = self.path.read_bytes()
        original_mtime = self.path.stat().st_mtime_ns
        self.enter("une note non enregistrée")
        self.app.search.set("école")
        self.app.category_filter.set("Mots")
        for language in ("fr", "en", "zh"):
            with self.subTest(language=language):
                self.app.set_language(language, persist=False)
                self.root.update()
                output = self.folder / f"export-{language}.docx"
                with patch("french_notes.ui.filedialog.asksaveasfilename", return_value=str(output)) as dialog:
                    with patch("french_notes.ui.messagebox.showinfo") as info:
                        self.app.export_all()
                self.assertTrue(output.exists())
                self.assertEqual(dialog.call_args.kwargs["title"], translate(language, "export_title"))
                self.assertEqual(dialog.call_args.kwargs["initialfile"], translate(language, "export_filename"))
                self.assertEqual(dialog.call_args.kwargs["filetypes"], [(translate(language, "word_files"), "*.docx")])
                document = Document(output)
                self.assertEqual(document.paragraphs[0].text, translate(language, "title"))
                self.assertEqual(
                    [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"],
                    [translate(language, "words"), translate(language, "sentences")],
                )
                text = "\n".join(paragraph.text for paragraph in document.paragraphs)
                self.assertIn("école", text)
                self.assertIn("Je suis étudiant.", text)
                self.assertNotIn("une note non enregistrée", text)
                self.assertEqual(info.call_args.args[0], translate(language, "export_done"))
                self.assertIn(translate(language, "export_summary", count=2), info.call_args.args[1])
                self.assertEqual(self.app.status.get(), translate(language, "export_summary", count=2))
                self.assertEqual(self.app.input.get("1.0", "end-1c"), "une note non enregistrée")
                self.assertEqual(self.path.read_bytes(), original_csv)
                self.assertEqual(self.path.stat().st_mtime_ns, original_mtime)
        self.error_dialog.assert_not_called()


if __name__ == "__main__":
    unittest.main()
